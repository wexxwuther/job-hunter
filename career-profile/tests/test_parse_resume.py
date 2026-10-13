"""Regression tests for parse_resume.py on real Word-authored resumes.

Fixtures in tests/fixtures/resumes/ are a FICTIONAL person's resume in five common
layouts, built with python-docx, then re-saved and PDF-exported by Microsoft Word
(2026-10-12), plus a rasterized image-only ("scanned") PDF. Each layout broke the
DOCX path in a different way before the fix (see docs/continuity/open-questions.md).
"""
import importlib.util
from pathlib import Path

import pytest

docx = pytest.importorskip("docx")
pytest.importorskip("pypdf")

SCRIPT = Path(__file__).resolve().parent.parent / "scripts" / "parse_resume.py"
_spec = importlib.util.spec_from_file_location("parse_resume", SCRIPT)
pr = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(pr)

FIX = Path(__file__).resolve().parent / "fixtures" / "resumes"
NAME, EMAIL, PHONE = "Jordan Q. Rivera", "jordan.rivera@example.com", "(555) 010-4477"
FACTS = [NAME, EMAIL, PHONE, "Northwind Logistics", "Contoso Health", "cut stockouts 18%",
         "120 dashboards", "30 analyst-hours", "B.S. Statistics, University of Arizona, 2018",
         "Power BI", "dbt", "Tableau Desktop Specialist (2020)"]


def _norm(text):
    return " ".join(text.split())


@pytest.mark.parametrize("name", [
    "01-single-column.docx", "02-contact-in-header.docx", "03-two-column-table.docx",
    "04-contact-in-textbox.docx", "05-experience-table-order.docx",
    "03-two-column-table.pdf", "05-experience-table-order.pdf",
])
def test_every_fact_is_recovered(name):
    text = _norm(pr.parse(FIX / name))
    missing = [f for f in FACTS if f not in text]
    assert not missing, f"{name} lost {missing}"


@pytest.mark.parametrize("name", ["02-contact-in-header.docx", "04-contact-in-textbox.docx"])
def test_contact_info_outside_the_body_comes_first(name):
    """Header and text-box contact details were silently dropped before the fix."""
    text = pr.parse(FIX / name)
    assert text.index(EMAIL) < text.index("Experience")
    assert text.count(NAME) == 1


def test_tables_keep_document_order_and_sections():
    """An experience table between prose sections used to be appended at the end,
    so the work history was filed under CERTIFICATIONS."""
    text = pr.parse(FIX / "05-experience-table-order.docx")
    norm = _norm(text)
    assert norm.index("EXPERIENCE") < norm.index("Northwind Logistics") < norm.index("EDUCATION")
    sections = pr._basic_sections(text)
    assert "Northwind Logistics" in sections["experience"]
    assert "Northwind" not in sections.get("certifications", "")


def test_merged_cells_are_read_once():
    """A horizontally merged cell used to repeat once per grid column (bullets x3)."""
    text = pr.parse(FIX / "05-experience-table-order.docx")
    assert text.count("cut stockouts 18%") == 1
    assert text.count("30 analyst-hours") == 1


def test_layout_table_headings_stay_on_their_own_line():
    """Two-column (sidebar + main) table: the EXPERIENCE heading used to be glued to
    the sidebar with ' | ', so no experience section was detected."""
    sections = pr._basic_sections(pr.parse(FIX / "03-two-column-table.docx"))
    assert "Northwind Logistics" in sections["experience"]
    assert {"skills", "education", "certifications"} <= set(sections)


def test_textbox_with_vml_fallback_is_not_duplicated(tmp_path):
    """Older Word writes a text box twice (DrawingML Choice + VML Fallback)."""
    from docx.oxml import parse_xml
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    box = (f'<w:txbxContent xmlns:w="{W}"><w:p><w:r><w:t>{NAME}</w:t></w:r></w:p></w:txbxContent>')
    run = parse_xml(
        f'<w:r xmlns:w="{W}" xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
        f'<mc:AlternateContent><mc:Choice Requires="wps">{box}</mc:Choice>'
        f'<mc:Fallback>{box}</mc:Fallback></mc:AlternateContent></w:r>')
    d = docx.Document()
    d.add_paragraph()._p.append(run)
    d.add_paragraph("Experience")
    path = tmp_path / "vml.docx"
    d.save(path)
    assert pr.parse(path).count(NAME) == 1


def test_scanned_pdf_fails_loudly(capsys):
    """An image-only PDF used to return empty text with exit 0."""
    with pytest.raises(ValueError, match="scanned"):
        pr.parse(FIX / "06-scanned-image-only.pdf")
    assert pr.main([str(FIX / "06-scanned-image-only.pdf")]) == 2
    assert "scanned" in capsys.readouterr().err


def test_empty_text_file_fails_loudly(tmp_path):
    p = tmp_path / "blank.txt"
    p.write_text("   \n", encoding="utf-8")
    with pytest.raises(ValueError, match="no text"):
        pr.parse(p)
